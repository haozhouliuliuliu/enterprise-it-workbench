"""生成演示用 PDF/DOCX/TXT 样本（一次性运行，产物提交仓库）。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "sample_docs"
OUT.mkdir(exist_ok=True)


def build_pdf() -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path = OUT / "IT Onboarding Handbook.pdf"
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setFont("STSong-Light", 11)
    lines = [
        "IT Onboarding Handbook (fictional)",
        "",
        "Section 1: Before Day One",
        "The hiring department must submit the device request at least 3 working days before the start date.",
        "A device request must include employee id, department, device model and expected delivery date.",
        "",
        "Section 2: Account Provisioning",
        "Accounts are created after the employee record is active. Production access requires a separate approval.",
        "",
        "Section 3: Device Handover",
        "The employee signs the device handover form. The asset owner is updated in the asset register.",
        "",
        "Section 4: Compliance Check",
        "All onboarding evidence is retained for audit. Draft supplier documents are not valid approval evidence.",
    ]
    y = 800
    for line in lines:
        pdf.drawString(50, y, line)
        y -= 18
        if y < 60:
            pdf.showPage()
            pdf.setFont("STSong-Light", 11)
            y = 800
    pdf.save()


def build_docx() -> None:
    document = Document()
    document.add_heading("供应商设备交付说明（草稿）", level=1)
    document.add_heading("交付范围", level=2)
    document.add_paragraph("供应商可在订单确认后直接向员工交付笔记本电脑，无需公司工单。")
    document.add_heading("交付周期", level=2)
    document.add_paragraph("标准机型交付周期为 7 个工作日，加急为 3 个工作日。")
    document.add_heading("特殊条款", level=2)
    document.add_paragraph("本草稿称交付无需审批，与公司现行制度冲突；本草稿尚未生效，不得作为审批依据。")
    document.add_paragraph("任何跳过工单、跳过设备申请的做法都不符合公司制度。")
    document.save(OUT / "供应商设备交付说明（草稿）.docx")


def build_txt() -> None:
    text = """资产预留与发放操作规范

一、预留规则
设备申请审批通过后，IT 可为申请人预留一台可用设备。
预留以资产编号为单位，同一台设备在同一时间只能被一个申请预留。
预留记录必须包含申请任务号、资产编号、预留时间和操作人。

二、发放流程
发放前核对员工编号、设备型号和领用确认单。
发放完成后，资产状态更新为 in_use，领用人为员工姓名。

三、超时释放
预留超过 3 个工作日未发放的，自动释放回可用库存。
释放操作必须记录审计事件，不得删除原预留记录。

四、并发要求
两个申请同时预留同一台设备时，只能有一个成功。
失败的一方必须重新选择可用设备，不得覆盖已有预留。
"""
    (OUT / "资产预留与发放操作规范.txt").write_text(text, encoding="utf-8")


if __name__ == "__main__":
    build_pdf()
    build_docx()
    build_txt()
    print("generated", list(OUT.iterdir()))
